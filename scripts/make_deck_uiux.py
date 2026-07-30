# -*- coding: utf-8 -*-
"""UI/UX刷新と効果の説明資料（クライアント向け）→ PowerPoint 生成"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

NAVY=RGBColor(0x1F,0x29,0x37); BLUE=RGBColor(0x1C,0xB0,0xF6); GREEN=RGBColor(0x10,0xB9,0x81)
ORANGE=RGBColor(0xFF,0x96,0x00); RED=RGBColor(0xEF,0x44,0x44); PURPLE=RGBColor(0x8B,0x5C,0xF6)
GRAY=RGBColor(0x6B,0x72,0x80); LIGHT=RGBColor(0xF3,0xF5,0xF8); WHITE=RGBColor(0xFF,0xFF,0xFF)
DARK=RGBColor(0x2B,0x2F,0x36); LINE=RGBColor(0xE0,0xE4,0xEA); MUTED=RGBColor(0x9A,0xA5,0xB1)

prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
BLANK=prs.slide_layouts[6]
def slide(): return prs.slides.add_slide(BLANK)
def box(s,x,y,w,h,fill=None,line=None,lw=1.0,round_=False):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE,
        Inches(x),Inches(y),Inches(w),Inches(h))
    if fill is None: sh.fill.background()
    else: sh.fill.solid(); sh.fill.fore_color.rgb=fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb=line; sh.line.width=Pt(lw)
    sh.shadow.inherit=False; return sh
def text(s,x,y,w,h,runs,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP,sa=4):
    tb=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=tb.text_frame
    tf.word_wrap=True; tf.vertical_anchor=anchor
    tf.margin_left=tf.margin_right=Pt(2); tf.margin_top=tf.margin_bottom=Pt(2)
    for i,para in enumerate(runs):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment=align; p.space_after=Pt(sa); p.space_before=Pt(0)
        for (t,sz,col,b) in para:
            r=p.add_run(); r.text=t; r.font.size=Pt(sz); r.font.color.rgb=col; r.font.bold=b; r.font.name="Meiryo"
    return tb
def header(s,kick,title,idx):
    box(s,0,0,13.333,1.15,fill=NAVY); box(s,0,1.15,13.333,0.06,fill=BLUE)
    text(s,0.6,0.16,11.5,0.45,[[(kick,12,BLUE,True)]])
    text(s,0.6,0.48,11.5,0.6,[[(title,23,WHITE,True)]])
    text(s,12.3,0.42,0.7,0.5,[[(f"{idx:02d}",18,RGBColor(0x55,0x60,0x70),True)]],align=PP_ALIGN.RIGHT)
def bullets(s,x,y,w,items,gap=0.44,dot=BLUE,size=12.5):
    for i,it in enumerate(items):
        yy=y+i*gap; box(s,x,yy+0.07,0.11,0.11,fill=dot,round_=True)
        if isinstance(it,tuple):
            text(s,x+0.26,yy-0.04,w-0.28,gap,[[(it[0],size,DARK,True),(it[1],size,GRAY,False)]])
        else:
            text(s,x+0.26,yy-0.04,w-0.28,gap,[[(it,size,DARK,False)]])

# ── 1 Title ─────────────────────────────────────────────
s=slide(); box(s,0,0,13.333,7.5,fill=NAVY); box(s,0,5.0,13.333,0.08,fill=BLUE)
text(s,0.9,1.7,11.5,1.0,[[("面談フィードバックツール　刷新のご報告",34,WHITE,True)]])
text(s,0.9,2.85,11.5,0.8,[[("「AIが点数をつける」から「話し方の傾向が見える」へ",18,BLUE,False)]])
text(s,0.9,5.25,11.5,0.5,[[("UI/UX の変更点と、期待できる効果",15,WHITE,True)]])
text(s,0.9,5.8,11.5,0.5,[[("2026.07  wheelsUp 面談FB",12,MUTED,False)]])

# ── 2 なぜ変えたか ──────────────────────────────────────
s=slide(); header(s,"WHY","なぜ方針を変えたのか",1)
box(s,0.6,1.5,6.0,2.9,fill=RGBColor(0xFE,0xF2,0xF2),round_=True); box(s,0.6,1.5,6.0,0.62,fill=RED,round_=True)
text(s,0.85,1.6,5.6,0.42,[[("これまでの課題（AI自動採点）",14,WHITE,True)]])
bullets(s,0.95,2.35,5.4,[
 "面談ごとの深い文脈理解をAIに求めていた",
 "点数の妥当性が現場感覚と噛み合わない",
 "処理が重く、途中で止まることがあった",
 "機能が多層化し「どこを見ればよいか」が不明瞭",
],gap=0.5,dot=RED,size=12.5)
box(s,6.9,1.5,5.8,2.9,fill=RGBColor(0xEC,0xFD,0xF5),round_=True); box(s,6.9,1.5,5.8,0.62,fill=GREEN,round_=True)
text(s,7.15,1.6,5.4,0.42,[[("新しい方針（トーク傾向の可視化）",14,WHITE,True)]])
bullets(s,7.25,2.35,5.2,[
 "良し悪しを判定せず「傾向」を示す",
 "計算だけで出すのでAI不要・待ち時間ほぼゼロ",
 "数字の根拠が明確（誰が見ても同じ結果）",
 "画面を絞り、見るべき場所を1つに",
],gap=0.5,dot=GREEN,size=12.5)
box(s,0.6,4.65,12.1,1.9,fill=LIGHT,round_=True)
text(s,0.9,4.8,11.5,0.4,[[("考え方の転換",13.5,NAVY,True)]])
text(s,0.9,5.3,11.6,1.1,[
 [("AIに「この面談は何点か」を判断させるのは、現時点では精度が安定しません。",12.5,DARK,False)],
 [("そこで「誰がどれだけ話し、どんな質問をしたか」という",12.5,DARK,False),("客観的な事実の可視化",12.5,GREEN,True),("に絞りました。",12.5,DARK,False)],
 [("評価は人（小林さん・西村さん）が行い、AIは判断材料の提示に徹する——という役割分担です。",12.5,GRAY,False)],
],sa=3)

# ── 3 UI Before / After ─────────────────────────────────
s=slide(); header(s,"UI CHANGE","画面はこう変わりました",2)
# BEFORE
box(s,0.6,1.45,6.0,5.3,fill=RGBColor(0xFA,0xFA,0xFB),round_=True,line=LINE)
box(s,0.6,1.45,6.0,0.55,fill=RGBColor(0x9C,0xA3,0xAF),round_=True)
text(s,0.85,1.54,5.6,0.4,[[("BEFORE",12,WHITE,True),("   階級型・機能過多",11,WHITE,False)]])
rows_b=[("リーダー画面 / メンバー画面が別々","役割ごとに別の画面。行き来ができない"),
("自分の面談しか見られない","他メンバーの面談は非表示"),
("1面談に12種類以上の情報が縦積み","5軸スコア・評価理由・軸別FB・タイムライン…"),
("フィードバックはリーダーのみ","現場同士の指摘が入らない"),
("採点待ちが発生（数十秒〜）","AI処理が終わるまで何も見えない")]
y=2.15
for (h,b) in rows_b:
    text(s,0.95,y,5.4,0.4,[[("✕ "+h,11.5,DARK,True)]])
    text(s,1.2,y+0.34,5.15,0.4,[[(b,10,GRAY,False)]])
    y+=0.92
# AFTER
box(s,6.9,1.45,5.8,5.3,fill=RGBColor(0xF0,0xFD,0xF9),round_=True,line=RGBColor(0xA7,0xF3,0xD0))
box(s,6.9,1.45,5.8,0.55,fill=GREEN,round_=True)
text(s,7.15,1.54,5.4,0.4,[[("AFTER",12,WHITE,True),("   フラット・要点のみ",11,WHITE,False)]])
rows_a=[("全員が同じ1画面","役割による画面の違いをなくした"),
("メンバータブで誰の面談でも閲覧","タブを押すだけで切り替わる"),
("見るのは「傾向」と「コメント」の2つ","不要な表示を撤去し縦の情報量を削減"),
("誰でもフィードバックを書ける","定型文ボタンでワンタップ入力"),
("開いた瞬間に結果が出る","計算のみなので待ち時間なし")]
y=2.15
for (h,b) in rows_a:
    text(s,7.25,y,5.2,0.4,[[("○ "+h,11.5,DARK,True)]])
    text(s,7.5,y+0.34,4.95,0.4,[[(b,10,GRAY,False)]])
    y+=0.92

# ── 4 画面イメージ ──────────────────────────────────────
s=slide(); header(s,"SCREEN","実際の画面構成",3)
text(s,0.6,1.35,12,0.4,[[("メンバータブを選ぶ → サマリーが出る → 個別の面談を開く、の3ステップだけです。",12.5,GRAY,False)]])
# タブ行
box(s,0.6,1.85,12.1,0.62,fill=WHITE,line=LINE,round_=True)
tabs=[("小林",PURPLE,True),("西村",BLUE,False),("辻内",GREEN,False),("安藤",ORANGE,False),("村上",RED,False)]
tx=0.85
for (nm,c,act) in tabs:
    box(s,tx,1.97,1.5,0.38,fill=c if act else WHITE,line=None if act else LINE,round_=True)
    text(s,tx,2.02,1.5,0.3,[[(nm,11,WHITE if act else GRAY,True)]],align=PP_ALIGN.CENTER)
    tx+=1.62
text(s,9.4,2.02,3.1,0.3,[[("← タブで担当者を切替",10.5,MUTED,False)]])
# サマリー
box(s,0.6,2.65,12.1,1.75,fill=RGBColor(0xF8,0xFA,0xFC),line=LINE,round_=True)
text(s,0.9,2.78,6,0.35,[[("① トーク傾向サマリー（その人の全面談を集計）",12,NAVY,True)]])
cards=[("会話スタイル","対話・深掘り型",PURPLE),("平均発話比率","52 %",BLUE),
("平均質問数","8.3 回",GREEN),("発話速度","310 字/分",ORANGE),("キャッチボール度","74 %",RED)]
cx=0.9
for (lb,vl,c) in cards:
    box(s,cx,3.2,2.25,1.05,fill=WHITE,line=LINE,round_=True)
    text(s,cx,3.32,2.25,0.3,[[(lb,9.5,GRAY,False)]],align=PP_ALIGN.CENTER)
    text(s,cx,3.62,2.25,0.4,[[(vl,13,c,True)]],align=PP_ALIGN.CENTER)
    cx+=2.4
# 面談カード
box(s,0.6,4.6,12.1,2.15,fill=RGBColor(0xF8,0xFA,0xFC),line=LINE,round_=True)
text(s,0.9,4.73,6,0.35,[[("② 個別の面談を開くと",12,NAVY,True)]])
det=[("📊 トーク傾向","発話比率／質問の質／言い切り度／相槌／口癖／前半・中盤・後半の推移",BLUE),
("💬 フィードバック","誰でも入力可。「深掘りが浅い」等の定型文をワンタップで追加",ORANGE),
("📋 候補者情報を整形","議事録を初回診断の5項目に整形（LARK提出形式・任意）",GREEN),
("📝 議事録","全文を確認。傾向の該当箇所をクリックすると本文にジャンプ",PURPLE)]
y=5.15
for (h,b,c) in det:
    box(s,0.95,y+0.04,0.1,0.28,fill=c)
    text(s,1.2,y,3.0,0.32,[[(h,11,DARK,True)]])
    text(s,4.3,y,8.2,0.32,[[(b,10.5,GRAY,False)]])
    y+=0.42

# ── 5 何が見えるようになったか ──────────────────────────
s=slide(); header(s,"WHAT YOU SEE","具体的に何が見えるようになったか",4)
items=[("誰がどれだけ話したか","発話比率（担当◯% / 相手◯%）。「話しすぎ」「聞けている」が数字で分かる",BLUE),
("質問の量と質","質問回数に加え、オープン質問（なぜ・どんな）とクローズド質問の内訳",GREEN),
("会話のリズム","話者交代の頻度、相槌の割合、一番長く話した場面の長さと中身",ORANGE),
("面談内の主導権の移り変わり","前半・中盤・後半の発話比率。「前半は聞き、後半に提案」等の流れ",PURPLE),
("言い切り度","断定表現と曖昧表現（〜かも／〜と思います）の比率",RED),
("口癖・頻出ワード","「なんか」「まあ」等の回数と、よく使う語彙",GRAY)]
y=1.5
for (h,b,c) in items:
    box(s,0.6,y,12.1,0.83,fill=WHITE,line=LINE,lw=1.1,round_=True); box(s,0.6,y,0.13,0.83,fill=c)
    text(s,0.95,y+0.09,3.6,0.35,[[(h,12.5,DARK,True)]])
    text(s,4.7,y+0.1,7.85,0.6,[[(b,11,GRAY,False)]],anchor=MSO_ANCHOR.MIDDLE)
    y+=0.9
box(s,0.6,y+0.05,12.1,0.72,fill=RGBColor(0xEA,0xF4,0xFE),round_=True)
text(s,0.9,y+0.13,11.5,0.6,[[("いずれも計算で求めた事実です。AIの判断は入らないため、",12,DARK,False),
 ("何度見ても同じ結果",12,BLUE,True),("になります。",12,DARK,False)]],anchor=MSO_ANCHOR.MIDDLE)

# ── 6 効果 ──────────────────────────────────────────────
s=slide(); header(s,"EFFECT","このシステムが生む効果",5)
eff=[("1","振り返りが「主観」から「事実」に",
 "点数の妥当性を議論する必要がなくなり、「今回は自分が7割話していた」という事実から会話が始まります。",GREEN),
("2","指摘が届きやすくなる",
 "誰でも他メンバーの面談にコメントできるため、指導がリーダー1人に集中しません。定型文ボタンで入力の手間も最小化。",BLUE),
("3","育成の型が共有される",
 "「良い面談」の傾向（発話比率・質問の質）がチームで可視化され、目指す状態を数字で共有できます。",ORANGE),
("4","データが蓄積し比較できる",
 "CSV出力により、担当者間の比較や、同じ人の時系列変化（成長の推移）を分析できます。",PURPLE)]
y=1.5
for (n,h,b,c) in eff:
    box(s,0.6,y,12.1,1.2,fill=LIGHT,round_=True); box(s,0.6,y,0.14,1.2,fill=c)
    box(s,0.95,y+0.28,0.62,0.62,fill=c,round_=True)
    text(s,0.95,y+0.34,0.62,0.5,[[(n,20,WHITE,True)]],align=PP_ALIGN.CENTER)
    text(s,1.85,y+0.18,10.5,0.4,[[(h,14.5,DARK,True)]])
    text(s,1.85,y+0.63,10.6,0.5,[[(b,11.5,GRAY,False)]])
    y+=1.3
box(s,0.6,y+0.05,12.1,0.85,fill=RGBColor(0xFF,0xF7,0xED),line=RGBColor(0xFE,0xD7,0xAA),round_=True)
text(s,0.9,y+0.14,11.5,0.7,[[("※ ",11,ORANGE,True),
 ("効果は運用を通じて検証していく前提です。まずは「傾向が見える状態」を作り、現場で使いながら指標の見方を調整していきます。",11,DARK,False)]],
 anchor=MSO_ANCHOR.MIDDLE)

# ── 7 使い方 ────────────────────────────────────────────
s=slide(); header(s,"HOW TO USE","日々の使い方",6)
steps=[("1","面談を実施","Google Meet で面談し、\n文字起こしを保存",BLUE),
("2","自動で取り込み","Driveに入った議事録を\nシステムが自動取得",GREEN),
("3","傾向を確認","タブを開くだけ。\n計算済みの傾向が表示",ORANGE),
("4","コメントする","気づいた点を入力。\n定型文ボタンも利用可",PURPLE),
("5","CSVで振り返り","月次でデータを出力し\n変化を確認",RED)]
x=0.55
for i,(n,h,b,c) in enumerate(steps):
    box(s,x,2.1,2.25,2.0,fill=LIGHT,round_=True); box(s,x,2.1,2.25,0.12,fill=c)
    box(s,x+0.85,2.35,0.55,0.55,fill=c,round_=True)
    text(s,x+0.85,2.4,0.55,0.45,[[(n,17,WHITE,True)]],align=PP_ALIGN.CENTER)
    text(s,x+0.15,3.02,1.95,0.35,[[(h,13,DARK,True)]],align=PP_ALIGN.CENTER)
    text(s,x+0.12,3.4,2.0,0.6,[[(b,10,GRAY,False)]],align=PP_ALIGN.CENTER)
    if i<4: text(s,x+2.22,2.85,0.35,0.4,[[("→",18,MUTED,True)]])
    x+=2.5
box(s,0.6,4.5,12.1,2.0,fill=RGBColor(0xEA,0xF4,0xFE),round_=True)
text(s,0.9,4.65,11.5,0.4,[[("運用のポイント",13,BLUE,True)]])
bullets(s,1.0,5.15,11.4,[
 ("日常の操作は「開いて見る」「コメントする」の2つだけ。","事前の準備や設定は不要です。"),
 ("新しい面談が入るとLarkに1回だけ通知が届きます。","（再確認しても通知は飛びません）"),
 ("議事録は Google Meet の自動文字起こし形式が最も精度よく分析できます。",""),
],gap=0.44,dot=BLUE,size=12)

# ── 8 まとめ ────────────────────────────────────────────
s=slide(); header(s,"SUMMARY","まとめ",7)
box(s,0.6,1.5,12.1,1.5,fill=NAVY,round_=True)
text(s,0.95,1.68,11.5,1.2,[
 [("AIに点数をつけさせるのをやめ、",15,WHITE,False),("「話し方の事実」を見える化",15,BLUE,True),("しました。",15,WHITE,False)],
 [("評価は人が行い、システムは判断材料を素早く正確に提示する役割に徹します。",13,MUTED,False)],
],sa=6)
cols=[("変わったこと",[
 "階級型UI → 全員共通のフラットな1画面",
 "自分の面談のみ → 誰の面談でも閲覧可",
 "12種類以上の情報 → 傾向とコメントの2つ",
 "リーダーのみFB → 誰でもFB入力可",
 "採点待ち数十秒 → 待ち時間なし"],BLUE),
("変わらないこと",[
 "議事録のDrive自動取り込み",
 "Larkへの通知（初回のみ）",
 "候補者情報の5項目整形（任意）",
 "議事録の全文確認",
 "データはこれまで通り蓄積"],GREEN)]
x=0.6
for (h,items,c) in cols:
    box(s,x,3.2,5.95,3.3,fill=LIGHT,round_=True); box(s,x,3.2,5.95,0.58,fill=c,round_=True)
    text(s,x+0.25,3.3,5.5,0.4,[[(h,13.5,WHITE,True)]])
    bullets(s,x+0.3,3.95,5.4,items,gap=0.47,dot=c,size=11.5)
    x+=6.15

out="/home/user/wheelsUp_Sales/docs/UIUX刷新と効果_クライアント説明.pptx"
prs.save(out)
print("SAVED:",out,"slides:",len(prs.slides._sldIdLst))
