# -*- coding: utf-8 -*-
"""wheelsUp 開発全体像（運用会社向け引き継ぎ資料）→ PowerPoint 生成"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

NAVY=RGBColor(0x1F,0x29,0x37); BLUE=RGBColor(0x1C,0xB0,0xF6); GREEN=RGBColor(0x10,0xB9,0x81)
ORANGE=RGBColor(0xFF,0x96,0x00); RED=RGBColor(0xFF,0x4B,0x4B); PURPLE=RGBColor(0x8B,0x5C,0xF6)
GRAY=RGBColor(0x6B,0x72,0x80); LIGHT=RGBColor(0xF3,0xF5,0xF8); WHITE=RGBColor(0xFF,0xFF,0xFF)
DARK=RGBColor(0x2B,0x2F,0x36); LINE=RGBColor(0xE0,0xE4,0xEA)

prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
BLANK=prs.slide_layouts[6]
def slide(): return prs.slides.add_slide(BLANK)
def box(s,x,y,w,h,fill=None,line=None,lw=1.0,round_=False):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE,
        Inches(x),Inches(y),Inches(w),Inches(h))
    if fill is None: sh.fill.background()
    else: sh.fill.solid(); sh.fill.fore_color.rgb=fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb=line; sh.line.width=Pt(lw)
    sh.shadow.inherit=False; return sh
def text(s,x,y,w,h,runs,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP,sa=4):
    tb=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=tb.text_frame
    tf.word_wrap=True; tf.vertical_anchor=anchor
    tf.margin_left=tf.margin_right=Pt(2); tf.margin_top=tf.margin_bottom=Pt(2)
    for i,para in enumerate(runs):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment=align; p.space_after=Pt(sa); p.space_before=Pt(0)
        for (t,sz,col,b) in para:
            r=p.add_run(); r.text=t; r.font.size=Pt(sz); r.font.color.rgb=col; r.font.bold=b; r.font.name="Meiryo"
    return tb
def header(s,kick,title,idx):
    box(s,0,0,13.333,1.15,fill=NAVY); box(s,0,1.15,13.333,0.06,fill=BLUE)
    text(s,0.6,0.16,11.5,0.45,[[(kick,12,BLUE,True)]])
    text(s,0.6,0.48,11.5,0.6,[[(title,23,WHITE,True)]])
    text(s,12.3,0.42,0.7,0.5,[[(f"{idx:02d}",18,RGBColor(0x55,0x60,0x70),True)]],align=PP_ALIGN.RIGHT)
def bullets(s,x,y,w,items,gap=0.46,dot=BLUE,size=13):
    for i,it in enumerate(items):
        yy=y+i*gap; box(s,x,yy+0.07,0.12,0.12,fill=dot,round_=True)
        if isinstance(it,tuple):
            text(s,x+0.28,yy-0.04,w-0.3,gap,[[(it[0],size,DARK,True),(it[1],size,GRAY,False)]])
        else:
            text(s,x+0.28,yy-0.04,w-0.3,gap,[[(it,size,DARK,False)]])

# 1 Title
s=slide(); box(s,0,0,13.333,7.5,fill=NAVY); box(s,0,5.0,13.333,0.08,fill=BLUE)
text(s,0.9,1.7,11.5,1.0,[[("wheelsUp  開発全体像",40,WHITE,True)]])
text(s,0.9,2.8,11.5,0.7,[[("営業面談フィードバック & RA開拓支援システム — 現状と運用引き継ぎ",17,BLUE,False)]])
text(s,0.9,5.25,11.5,0.5,[[("運用会社向け 引き継ぎ資料",16,WHITE,True)]])
text(s,0.9,5.8,11.5,0.5,[[("2026.07  /  最新: トーク傾向分析へピボット中",12,RGBColor(0x9A,0xA5,0xB1),False)]])

# 2 サマリー
s=slide(); header(s,"SUMMARY","システムの全体像（1枚まとめ）",1)
cards=[(0.6,BLUE,"何のシステムか","営業(キャリアコンサル)の面談議事録を分析し、育成に使うWebアプリ。2機能=①面談FB ②RA開拓(候補者/企業の発掘支援)。"),
(2.35,GREEN,"現在の方針（重要）","AI自動採点は実運用が難しく、『話し方の癖・傾向を機械的に可視化』する方向へピボット中(2026/7 決定)。"),
(4.1,ORANGE,"技術構成","React+Vite(フロント) / Vercel Serverless(API) / Supabase(DB) / Google Gemini(AI)。外部連携=Drive・Lark・Pipedrive。"),
(5.85,PURPLE,"運用状態","本番稼働中(Vercel)。議事録はGoogle Drive経由で自動取込→分析→表示。デプロイはCLIから手動。")]
for (y,col,h,b) in cards:
    box(s,0.6,y,12.1,1.6,fill=LIGHT,round_=True); box(s,0.6,y,0.14,1.6,fill=col)
    text(s,0.95,y+0.16,11.5,0.5,[[(h,15,DARK,True)]]); text(s,0.95,y+0.66,11.6,0.85,[[(b,12.5,GRAY,False)]])

# 3 全体構成（2機能）
s=slide(); header(s,"SYSTEM MAP","2つの機能で構成される",2)
# 面談FB
box(s,0.6,1.5,6.0,5.2,fill=RGBColor(0xEA,0xF4,0xFE),round_=True); box(s,0.6,1.5,6.0,0.7,fill=BLUE,round_=True)
text(s,0.85,1.62,5.6,0.45,[[("① 面談FB（メインの育成機能）",15,WHITE,True)]])
bullets(s,0.95,2.5,5.3,[
 "議事録の取込(Drive自動連携 / 手動貼付)",
 "話者分離(「あなた」=担当者本人)",
 ("トーク傾向分析 ","発話比率/質問数/口癖/速度(機械計算)"),
 "候補者情報の整形(初回診断5項目)",
 "AI 5軸採点(縮小方針・任意)",
 "フィードバック/校正/手動タグ",
 "学習ライブラリ(型+実例トーク)",
 "CSVエクスポート(振り返り用)",
],gap=0.5,dot=BLUE,size=12)
# RA開拓
box(s,6.9,1.5,5.8,5.2,fill=RGBColor(0xEA,0xF7,0xF0),round_=True); box(s,6.9,1.5,5.8,0.7,fill=GREEN,round_=True)
text(s,7.15,1.62,5.4,0.45,[[("② RA開拓（候補者/企業の発掘）",15,WHITE,True)]])
bullets(s,7.25,2.5,5.1,[
 "企業/候補者データの投入・管理",
 "求人ページのクロール(自動収集)",
 "候補者×求人のマッチング採点",
 "新規企業のAI発掘(Gemini)",
 "URL自動補完・再検証",
 "全自動収集モード(定期ループ)",
 "Pipedrive連携",
],gap=0.5,dot=GREEN,size=12)

# 4 技術スタック
s=slide(); header(s,"ARCHITECTURE","技術スタックとアーキテクチャ",3)
layers=[("フロントエンド","React + Vite + TypeScript / TailwindCSS / React Query / React Router","ブラウザで動くSPA。面談FBとRA開拓を1画面で切替",BLUE),
("APIサーバ","Vercel Serverless Functions (Node.js/TypeScript)","/api/meetings, /api/ra ほか。1関数=最大60秒の制約あり",ORANGE),
("データベース","Supabase (PostgreSQL)","面談議事録・スコア・校正・企業/候補者データを保存。APIはservice_roleでアクセス",GREEN),
("AI","Google Gemini (2.5-flash / 2.0-flash / flash-lite の多段フォールバック)","整形・採点・発掘に利用。OpenAIも接続可",PURPLE),
("外部連携","Google Drive(SA) / Mimo(Apps Script) / Lark・Slack(通知) / Pipedrive","議事録の自動取込、採点通知、CRM連携",RED)]
y=1.45
for (h,tech,desc,col) in layers:
    box(s,0.6,y,12.1,1.0,fill=WHITE,line=LINE,lw=1.2,round_=True); box(s,0.6,y,0.14,1.0,fill=col)
    text(s,0.95,y+0.1,2.6,0.8,[[(h,13.5,DARK,True)]],anchor=MSO_ANCHOR.MIDDLE)
    text(s,3.5,y+0.12,9.0,0.4,[[(tech,12,DARK,True)]])
    text(s,3.5,y+0.52,9.0,0.4,[[(desc,11,GRAY,False)]])
    y+=1.06

# 5 データフロー
s=slide(); header(s,"DATA FLOW","運用フロー：議事録が入ってから表示まで",4)
steps=[("① 面談",  "Google Meetで面談\n→自動文字起こし",BLUE),
("② 保存","Mimo拡張がDriveに\n議事録を格納",GREEN),
("③ 取込","APIがDriveから取込\n(Webhook/定期)",ORANGE),
("④ 分析","話者分離+トーク傾向分析\n(+任意でAI整形/採点)",PURPLE),
("⑤ 表示・通知","画面に傾向表示\n初回はLarkへ1通知",RED)]
x=0.55
for i,(h,b,col) in enumerate(steps):
    box(s,x,2.3,2.25,1.8,fill=LIGHT,round_=True); box(s,x,2.3,2.25,0.12,fill=col)
    text(s,x+0.15,2.55,1.95,0.5,[[(h,14,DARK,True)]])
    text(s,x+0.15,3.05,1.95,0.9,[[(b,10.5,GRAY,False)]])
    if i<4: text(s,x+2.22,2.75,0.35,0.5,[[("→",20,GRAY,True)]])
    x+=2.5
box(s,0.6,4.6,12.1,1.9,fill=RGBColor(0xEA,0xF4,0xFE),round_=True)
text(s,0.9,4.75,11.5,0.4,[[("運用ポイント",13,BLUE,True)]])
bullets(s,1.0,5.25,11.4,[
 ("自動化済み：","Drive追加→自動で取込・分析・(初回のみ)Lark通知。担当者は基本ボタン操作不要"),
 ("手動操作：","再分析・整形・採点・フィードバック入力は画面のボタンから任意で実行"),
 ("通知ルール：","初回採点/分析のみLark通知。再実行では通知しない(連打防止)"),
],gap=0.42,dot=BLUE,size=12)

# 6 面談FBの現状（ピボット）
s=slide(); header(s,"PIVOT","面談FBの現状：採点→トーク傾向分析へ",5)
box(s,0.6,1.5,6.0,2.4,fill=RGBColor(0xFF,0xF0,0xF0),round_=True); box(s,0.6,1.5,6.0,0.6,fill=RED,round_=True)
text(s,0.85,1.6,5.6,0.4,[[("課題だった旧方針：AI自動採点",13.5,WHITE,True)]])
bullets(s,0.95,2.35,5.3,[
 "面談ごとに深い文脈理解をAIに要求→不安定",
 "点数の妥当性が現場感覚と合わない",
 "処理が重くタイムアウトしやすい",
],gap=0.46,dot=RED,size=12)
box(s,6.9,1.5,5.8,2.4,fill=RGBColor(0xEA,0xF7,0xF0),round_=True); box(s,6.9,1.5,5.8,0.6,fill=GREEN,round_=True)
text(s,7.15,1.6,5.4,0.4,[[("新方針：話し方の傾向を可視化",13.5,WHITE,True)]])
bullets(s,7.25,2.35,5.1,[
 "良し悪しを判定せず『傾向』を出す",
 "機械計算なのでAI不要・一瞬・安定",
 "発話比率/質問数/口癖/速度/頻出ワード",
],gap=0.46,dot=GREEN,size=12)
box(s,0.6,4.1,12.1,2.35,fill=LIGHT,round_=True)
text(s,0.9,4.25,11.5,0.4,[[("トーク傾向分析で出す指標（AI不要・実装済み）",13,DARK,True)]])
metrics=["発話比率(担当:相手)","質問数・質問率","平均発話文字数","最長独話","発話速度(字/分)","口癖トップ","頻出ワード","総ターン数"]
mx=0.9; my=4.8
for i,m in enumerate(metrics):
    col=[BLUE,GREEN,ORANGE,PURPLE][i%4]
    box(s,mx+(i%4)*2.95,my+(i//4)*0.72,2.8,0.6,fill=WHITE,line=LINE,round_=True)
    text(s,mx+(i%4)*2.95,my+(i//4)*0.72+0.13,2.8,0.4,[[(m,11.5,DARK,True)]],align=PP_ALIGN.CENTER)

# 7 主要API
s=slide(); header(s,"API","主要APIエンドポイント（/api/meetings ほか）",6)
groups=[("取込・連携",["drive-import / drive-webhook (Drive取込)","push-transcript (Apps Scriptからpush)","transcribe (音声→文字起こし)"],BLUE),
("分析・採点",[":id/diagnose (5項目整形)",":id/score (AI採点)","bulk-rescore / bulk-diagnose (一括)"],ORANGE),
("学習・校正",[":id/calibrate (良い/悪いマーク)",":id/manual-observation (手動タグ)","extract-style / training-health"],GREEN),
("その他",[":id/outcome (成果記録)",":id/leader-feedback (コメント)","/api/ra (RA開拓一式)"],PURPLE)]
x=0.6
for (h,items,col) in groups:
    box(s,x,1.5,2.95,4.8,fill=LIGHT,round_=True); box(s,x,1.5,2.95,0.55,fill=col,round_=True)
    text(s,x+0.15,1.6,2.7,0.4,[[(h,13,WHITE,True)]])
    for i,it in enumerate(items):
        text(s,x+0.2,2.25+i*0.95,2.6,0.9,[[("・"+it,10.5,DARK,False)]])
    x+=3.05
text(s,0.6,6.45,12,0.4,[[("※ /api/meetings は path で分岐する単一関数。合計25以上のサブ機能を集約。",10.5,GRAY,False)]])

# 8 環境変数・連携
s=slide(); header(s,"CONFIG","環境変数・外部サービス（運用に必要）",7)
cfg=[("必須(コア)","SUPABASE_URL / SUPABASE_SERVICE_ROLE_KEY","DB接続。これが無いと全機能停止",RED),
("AI","GEMINI_API_KEY (GEMINI_MODEL / GEMINI_SCORING_MODEL)","整形・採点・発掘。無料枠のRPM制限に注意",PURPLE),
("Drive連携","GOOGLE_SA_EMAIL / GOOGLE_SA_PRIVATE_KEY / MIMO_DRIVE_FOLDER_ID","議事録の自動取込(サービスアカウント)",BLUE),
("通知","LARK_WEBHOOK_URL / LARK_USER_IDS / SLACK_WEBHOOK_URL","採点/分析完了の通知。未設定なら通知OFF",GREEN),
("保護・その他","CRON_SECRET / DRIVE_WEBHOOK_TOKEN / APP_BASE_URL","Webhook認証・URL生成",ORANGE),
("RA開拓","GOOGLE_SEARCH_API_KEY / GOOGLE_SEARCH_CX / JINA_API_KEY / PIPEDRIVE_API_TOKEN","企業発掘・クロール・CRM連携",GRAY)]
y=1.5
for (h,keys,desc,col) in cfg:
    box(s,0.6,y,12.1,0.82,fill=WHITE,line=LINE,round_=True); box(s,0.6,y,0.14,0.82,fill=col)
    text(s,0.95,y+0.08,2.4,0.7,[[(h,12.5,DARK,True)]],anchor=MSO_ANCHOR.MIDDLE)
    text(s,3.3,y+0.06,9.3,0.4,[[(keys,11,DARK,True)]])
    text(s,3.3,y+0.44,9.3,0.35,[[(desc,10,GRAY,False)]])
    y+=0.88
text(s,0.6,6.9,12,0.4,[[("※ 値は Vercel の Environment Variables で管理。本資料にはキー名のみ記載(秘匿値は非掲載)。",10,GRAY,False)]])

# 9 デプロイ・運用手順
s=slide(); header(s,"OPERATIONS","デプロイ・運用手順",8)
box(s,0.6,1.5,12.1,2.5,fill=RGBColor(0x1E,0x28,0x36),round_=True)
text(s,0.9,1.65,11.5,0.4,[[("デプロイ（本番反映）",13,BLUE,True)]])
text(s,0.9,2.15,11.6,1.7,[
 [("# リポジトリのルートで実行（wheelup-vercel の中では実行しない）",11,RGBColor(0x8A,0x94,0xA0),False)],
 [("cd ~/wheelsUp_Sales",13,WHITE,True)],
 [("git pull origin <branch>",13,WHITE,True)],
 [("cd wheelup-vercel && npm install   # 依存更新時のみ",12,RGBColor(0xB8,0xC2,0xCC),False)],
 [("cd ~/wheelsUp_Sales && vercel --prod",13,GREEN,True)],
],sa=3)
box(s,0.6,4.2,5.9,2.35,fill=LIGHT,round_=True)
text(s,0.85,4.33,5.5,0.4,[[("定常運用（自動）",12.5,GREEN,True)]])
bullets(s,0.95,4.85,5.3,[
 "議事録はDrive追加で自動取込・分析",
 "初回のみLark通知",
 "担当者は画面で傾向/FBを確認",
],gap=0.44,dot=GREEN,size=11.5)
box(s,6.7,4.2,6.0,2.35,fill=LIGHT,round_=True)
text(s,6.95,4.33,5.6,0.4,[[("メンテ観点",12.5,ORANGE,True)]])
bullets(s,7.05,4.85,5.4,[
 "Gemini APIキーの残枠・課金",
 "Driveサービスアカウントの共有権限",
 "一括処理は60秒制限を超えぬようバッチ実行",
],gap=0.44,dot=ORANGE,size=11.5)

# 10 制約・注意点
s=slide(); header(s,"CONSTRAINTS","既知の制約・注意点",9)
items=[("Vercelの関数は60秒上限","大量の一括処理(整形/採点)はタイムアウトする。→3件/1件ずつのバッチ+フロントでループする設計に。リアルタイム音声等は別基盤が必要。",RED),
("AI採点の精度に限界","文脈依存の点数化は不安定。→機械計算のトーク傾向分析を主軸に据えるピボットを実施済み。",ORANGE),
("Gemini無料枠のRPM制限","連続呼び出しは待機を挟む。件数が多い処理は時間がかかる。",PURPLE),
("デプロイはCLI手動","VercelのRoot Directoryが wheelup-vercel のため、必ずリポジトリのルートから vercel --prod を実行する。",BLUE),
("機能過多になりやすい","採点系UIが多層化していた。運用会社向けにはトーク傾向+CSV+FBのシンプル構成を推奨。",GREEN)]
y=1.5
for (h,b,col) in items:
    box(s,0.6,y,12.1,1.0,fill=LIGHT,round_=True); box(s,0.6,y,0.14,1.0,fill=col)
    text(s,0.95,y+0.12,11.5,0.4,[[(h,13.5,DARK,True)]])
    text(s,0.95,y+0.55,11.6,0.4,[[(b,11,GRAY,False)]])
    y+=1.06

# 11 これまでとロードマップ
s=slide(); header(s,"ROADMAP","これまでの経緯と今後",10)
text(s,0.6,1.4,12,0.4,[[("開発の流れ",14,NAVY,True)]])
timeline=[("〜2026/6","5軸AI採点+学習データ比較を構築。精度・UIで試行錯誤",BLUE),
("2026/6","話者分離・整形(01フォーマット)・小林データ比較を実装",GREEN),
("2026/7","AI採点の限界を判断→トーク傾向分析へピボット決定",ORANGE),
("現在","トーク傾向分析+CSV出力を実装。UIフラット化を予定",RED)]
y=1.9
for (d,b,col) in timeline:
    box(s,0.6,y,1.7,0.7,fill=col,round_=True); text(s,0.6,y+0.16,1.7,0.4,[[(d,11.5,WHITE,True)]],align=PP_ALIGN.CENTER)
    box(s,2.5,y,10.2,0.7,fill=LIGHT,round_=True); text(s,2.75,y+0.16,9.8,0.4,[[(b,12,DARK,False)]],anchor=MSO_ANCHOR.MIDDLE)
    y+=0.82
box(s,0.6,5.4,12.1,1.5,fill=RGBColor(0xEA,0xF4,0xFE),round_=True)
text(s,0.9,5.55,11.5,0.4,[[("今後の予定",13,BLUE,True)]])
bullets(s,1.0,6.05,11.4,[
 ("UIフラット化：","誰でも全員の面談を閲覧・FB入力できるタブ構造へ"),
 ("運用移管：","運用会社が自分たちで更新継続できるパッケージ化"),
],gap=0.42,dot=BLUE,size=12)

out="/home/user/wheelsUp_Sales/docs/wheelsUp開発全体像_運用会社向け.pptx"
prs.save(out)
print("SAVED:",out,"slides:",len(prs.slides._sldIdLst))
