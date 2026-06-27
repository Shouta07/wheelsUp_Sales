# -*- coding: utf-8 -*-
"""接客AI採点・ロープレ事業 提案資料 → PowerPoint 生成"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ---- カラーパレット ----
NAVY   = RGBColor(0x1F, 0x29, 0x37)
BLUE   = RGBColor(0x1C, 0xB0, 0xF6)
GREEN  = RGBColor(0x10, 0xB9, 0x81)
ORANGE = RGBColor(0xFF, 0x96, 0x00)
RED    = RGBColor(0xFF, 0x4B, 0x4B)
GRAY   = RGBColor(0x6B, 0x72, 0x80)
LIGHT  = RGBColor(0xF3, 0xF5, 0xF8)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
DARK   = RGBColor(0x2B, 0x2F, 0x36)

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]

def slide():
    return prs.slides.add_slide(BLANK)

def box(s, x, y, w, h, fill=None, line=None, line_w=1.0, round_=False):
    from pptx.enum.shapes import MSO_SHAPE
    shp = s.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    return shp

def text(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space_after=4):
    """runs: list of paragraphs; each paragraph = list of (txt,size,color,bold)"""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Pt(2); tf.margin_top = tf.margin_bottom = Pt(2)
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(space_after); p.space_before = Pt(0)
        for (t, sz, col, bold) in para:
            r = p.add_run(); r.text = t
            r.font.size = Pt(sz); r.font.color.rgb = col; r.font.bold = bold
            r.font.name = "Meiryo"
    return tb

def header(s, kicker, title, idx):
    box(s, 0, 0, 13.333, 1.15, fill=NAVY)
    box(s, 0, 1.15, 13.333, 0.06, fill=BLUE)
    text(s, 0.6, 0.16, 11.5, 0.45, [[(kicker, 12, BLUE, True)]])
    text(s, 0.6, 0.48, 11.5, 0.6, [[(title, 24, WHITE, True)]])
    text(s, 12.3, 0.42, 0.7, 0.5, [[(f"{idx:02d}", 18, RGBColor(0x55,0x60,0x70), True)]], align=PP_ALIGN.RIGHT)

def bullet(s, x, y, w, items, gap=0.46, dot=BLUE, size=13):
    for i, it in enumerate(items):
        yy = y + i*gap
        box(s, x, yy+0.07, 0.12, 0.12, fill=dot, round_=True)
        if isinstance(it, tuple):
            head, body = it
            text(s, x+0.28, yy-0.04, w-0.3, gap,
                 [[(head, size, DARK, True), (body, size, GRAY, False)]])
        else:
            text(s, x+0.28, yy-0.04, w-0.3, gap, [[(it, size, DARK, False)]])

# ============================================================
# Slide 1 — Title
# ============================================================
s = slide()
box(s, 0, 0, 13.333, 7.5, fill=NAVY)
box(s, 0, 5.0, 13.333, 0.08, fill=BLUE)
text(s, 0.9, 1.9, 11.5, 1.0, [[("接客スキル AI 採点・ロープレ事業", 40, WHITE, True)]])
text(s, 0.9, 3.0, 11.5, 0.7, [[("— 面談採点AI(wheelsUp)のエンジンを現場接客の育成へ転用する —", 18, BLUE, False)]])
text(s, 0.9, 5.3, 11.5, 0.5, [[("バイタリティデザイン合同会社 向け 事業提案", 16, WHITE, True)]])
text(s, 0.9, 5.85, 11.5, 0.5, [[("2026.06.27  /  ベース実装: wheelsUp Sales", 12, RGBColor(0x9A,0xA5,0xB1), False)]])

# ============================================================
# Slide 2 — エグゼクティブサマリー
# ============================================================
s = slide()
header(s, "EXECUTIVE SUMMARY", "結論：接客採点は転用可、ロープレは新規開発", 1)
cards = [
    (0.6, GREEN, "仮説は市場で実証され始めている",
     "IDOM・プリモGHD・フロンティアが2025〜26年にAI接客ロープレを導入。早期戦力化・指導負担減で成果（日経 2026/6/25）。"),
    (2.55, BLUE, "既存エンジンがそのまま効く（フェーズ1）",
     "「発話→観察抽出→ルーブリック採点→お手本比較」の構造は業界非依存。接客の実データ採点は業界差し替えでほぼ実現可能。"),
    (4.5, ORANGE, "AI客役ロープレは新規開発（フェーズ2）",
     "リアルタイム会話・音声/表情解析・常駐基盤が必要。今のVercel(60秒制限)では不可。別プロジェクトとして設計。"),
    (6.45, RED, "推奨：小さく早く実証する",
     "フェーズ1(接客採点PoC)を『1業界・1チェックリスト・1お手本』の別アプリで立ち上げ、市場の追い風を捉える。"),
]
for (y, col, head, body) in cards:
    box(s, 0.6, y, 12.1, 1.75, fill=LIGHT, round_=True)
    box(s, 0.6, y, 0.14, 1.75, fill=col)
    text(s, 0.95, y+0.18, 11.5, 0.5, [[(head, 16, DARK, True)]])
    text(s, 0.95, y+0.72, 11.6, 0.9, [[(body, 13, GRAY, False)]])

# ============================================================
# Slide 3 — 市場背景（ファクト）
# ============================================================
s = slide()
header(s, "MARKET", "市場背景：AI接客ロープレ導入が進む（日経 2026/6/25）", 2)
rows = [
    ("IDOM（ガリバー）", "新卒営業の接客ロープレ。話す速さ・声色・表情・必須ワード漏れをAIが審査", "25年入社が前年の利益水準を3ヶ月早く達成／店長の負担減", GREEN),
    ("プリモGHD", "画面上のAI顧客に指輪を売る接客ロープレ。自社商品・店舗運用に基づき開発", "接客の高レベル平準化＋指導役の負担軽減", BLUE),
    ("フロンティア", "テレアポ新人研修。30顧客像×シナリオ。約30項目を5段階評価", "営業人員を現場に残したまま新人を早期戦力化", ORANGE),
]
y = 1.5
box(s, 0.6, y, 3.0, 0.5, fill=NAVY); text(s,0.7,y+0.08,2.9,0.4,[[("企業",12,WHITE,True)]])
box(s, 3.6, y, 5.2, 0.5, fill=NAVY); text(s,3.7,y+0.08,5.1,0.4,[[("AI活用の内容",12,WHITE,True)]])
box(s, 8.8, y, 3.9, 0.5, fill=NAVY); text(s,8.9,y+0.08,3.8,0.4,[[("成果",12,WHITE,True)]])
y += 0.5
for (co, what, res, col) in rows:
    box(s, 0.6, y, 3.0, 1.3, fill=LIGHT); box(s,0.6,y,0.1,1.3,fill=col)
    box(s, 3.6, y, 5.2, 1.3, fill=WHITE, line=RGBColor(0xE0,0xE4,0xEA))
    box(s, 8.8, y, 3.9, 1.3, fill=WHITE, line=RGBColor(0xE0,0xE4,0xEA))
    text(s,0.85,y+0.12,2.7,1.1,[[(co,13,DARK,True)]],anchor=MSO_ANCHOR.MIDDLE)
    text(s,3.8,y+0.12,4.9,1.1,[[(what,11.5,GRAY,False)]],anchor=MSO_ANCHOR.MIDDLE)
    text(s,9.0,y+0.12,3.6,1.1,[[(res,11.5,DARK,False)]],anchor=MSO_ANCHOR.MIDDLE)
    y += 1.34
box(s, 0.6, y+0.05, 12.1, 0.85, fill=RGBColor(0xEAF,0xF3,0xFF) if False else RGBColor(0xEA,0xF4,0xFE), round_=True)
text(s, 0.85, y+0.13, 11.7, 0.7,
     [[("勝ちパターン：", 13, BLUE, True),
       ("①採点基準が具体的（必須ワード/チェックリスト） ②全店で標準化 ③指導側の負担減 ④KPI早期達成", 13, DARK, False)]],
     anchor=MSO_ANCHOR.MIDDLE)

# ============================================================
# Slide 4 — 既存エンジンの中核構造
# ============================================================
s = slide()
header(s, "CORE ENGINE", "既存システム(wheelsUp)の中核は業界非依存", 3)
text(s, 0.6, 1.4, 12, 0.4, [[("発話データを採点する汎用エンジン。4ステップは接客にもそのまま適用できる。", 13, GRAY, False)]])
steps = [
    ("① 観察抽出", "発話から観察可能な行動を引用付きで抽出", BLUE),
    ("② ルーブリック採点", "各項目の達成条件を満たしたかで点数化", GREEN),
    ("③ お手本比較", "トップ人材の基準・実データと対比", ORANGE),
    ("④ 構造化整形", "必須項目の充足/欠落を可視化", RED),
]
x = 0.6
for (h, b, col) in steps:
    box(s, x, 2.0, 2.9, 1.5, fill=LIGHT, round_=True); box(s, x, 2.0, 2.9, 0.12, fill=col)
    text(s, x+0.2, 2.25, 2.6, 0.5, [[(h, 14, DARK, True)]])
    text(s, x+0.2, 2.75, 2.6, 0.7, [[(b, 11, GRAY, False)]])
    if x < 9.5:
        text(s, x+2.78, 2.45, 0.4, 0.5, [[("→", 20, GRAY, True)]])
    x += 3.05
# mapping table
text(s, 0.6, 3.8, 12, 0.4, [[("接客への置き換え（そのまま転用）", 15, NAVY, True)]])
maps = [
    ("議事録 / 文字起こし", "接客の録音・文字起こし"),
    ("5軸（ニーズ/提案/信頼/前進/情報）", "接客チェックリスト（挨拶/ヒアリング/提案/CL/必須ワード）"),
    ("トップ人材(小林)のお手本データ", "トップ販売員の接客データ"),
    ("01診断「確認できず」検出", "必須ワード・確認事項の漏れ検出"),
]
y = 4.3
box(s, 0.6, y, 5.9, 0.46, fill=NAVY); text(s,0.75,y+0.07,5.7,0.4,[[("wheelsUp（転職面談）",12,WHITE,True)]])
box(s, 6.8, y, 5.9, 0.46, fill=GREEN); text(s,6.95,y+0.07,5.7,0.4,[[("接客版（バイタリティデザイン）",12,WHITE,True)]])
y += 0.46
for i,(a,b) in enumerate(maps):
    bg = WHITE if i%2==0 else LIGHT
    box(s, 0.6, y, 5.9, 0.5, fill=bg, line=RGBColor(0xE5,0xE8,0xED))
    box(s, 6.8, y, 5.9, 0.5, fill=bg, line=RGBColor(0xE5,0xE8,0xED))
    text(s,0.75,y+0.06,5.7,0.4,[[(a,11.5,GRAY,False)]],anchor=MSO_ANCHOR.MIDDLE)
    text(s,6.95,y+0.06,5.7,0.4,[[(b,11.5,DARK,True)]],anchor=MSO_ANCHOR.MIDDLE)
    y += 0.5

# ============================================================
# Slide 5 — 転用可能性マッピング
# ============================================================
s = slide()
header(s, "FEASIBILITY", "そのまま使える資産 / 作り替えが要る部分", 4)
box(s, 0.6, 1.45, 6.0, 5.4, fill=RGBColor(0xEA,0xF7,0xF0), round_=True)
box(s, 0.6, 1.45, 6.0, 0.7, fill=GREEN, round_=True)
text(s, 0.85, 1.6, 5.6, 0.45, [[("✅ そのまま使える（フェーズ1で即活用）", 15, WHITE, True)]])
bullet(s, 0.95, 2.45, 5.4, [
    "ルーブリック/チェックリスト採点エンジン",
    "観察抽出＋引用照合（捏造引用を弾く）",
    "お手本（トップ人材）との比較採点",
    "構造化整形（必須項目の充足/欠落の可視化）",
    "話者分離（スタッフ vs 客 の分離）",
    "軸別フィードバック＋学習ライブラリ",
    "採点モデル(ルーブリック)の現場可視化",
], gap=0.58, dot=GREEN, size=12.5)

box(s, 6.9, 1.45, 5.8, 5.4, fill=RGBColor(0xFF,0xF3,0xE6), round_=True)
box(s, 6.9, 1.45, 5.8, 0.7, fill=ORANGE, round_=True)
text(s, 7.15, 1.6, 5.4, 0.45, [[("⚠️ 作り替え/追加が必要（フェーズ2）", 15, WHITE, True)]])
bullet(s, 7.25, 2.45, 5.2, [
    "リアルタイムAI客役（会話するロープレ相手）",
    "音声解析（話す速さ・声色・間・抑揚）",
    "表情/映像解析（IDOMが審査）",
    "30項目×5段階の細粒度チェックリスト",
    "業界テンプレートの切り替え機構",
    "ストリーミング/常駐の実行基盤",
], gap=0.62, dot=ORANGE, size=12.5)

# ============================================================
# Slide 6 — プロダクト2段階構想
# ============================================================
s = slide()
header(s, "ROADMAP", "プロダクト2段階構想", 5)
# Phase1
box(s, 0.6, 1.5, 12.1, 2.45, fill=LIGHT, round_=True); box(s,0.6,1.5,0.16,2.45,fill=BLUE)
text(s, 0.95, 1.65, 11, 0.5, [[("フェーズ1：接客「実データ採点」", 18, BLUE, True), ("  — 短期・既存資産で勝てる", 13, GRAY, False)]])
text(s, 0.95, 2.25, 11.6, 0.5, [[("実際の接客（録音/文字起こし）をチェックリストで採点しフィードバック", 13, DARK, True)]])
bullet(s, 1.0, 2.75, 11.4, [
    "接客チェックリスト定義（必須ワード・接客フロー・NG行動）",
    "トップ販売員の接客を『お手本データ』に登録（小林ポジション）／既存コードの業界差し替えで最短実装",
], gap=0.42, dot=BLUE, size=12.5)
text(s, 0.95, 3.55, 11.6, 0.4, [[("価値：", 12.5, BLUE, True),("全店の接客品質の見える化・標準化、店長/トレーナーの採点負担減", 12.5, GRAY, False)]])
# Phase2
box(s, 0.6, 4.15, 12.1, 2.6, fill=LIGHT, round_=True); box(s,0.6,4.15,0.16,2.6,fill=ORANGE)
text(s, 0.95, 4.3, 11, 0.5, [[("フェーズ2：AI客役「リアルタイムロープレ」", 18, ORANGE, True), ("  — 中期・本命", 13, GRAY, False)]])
text(s, 0.95, 4.9, 11.6, 0.5, [[("AIが客として応答し、スタッフが練習。終了後に採点", 13, DARK, True)]])
bullet(s, 1.0, 5.4, 11.4, [
    "会話するAI客（ペルソナ×シナリオ／フロンティアの30顧客像が参考）",
    "リアルタイム音声対話（speed/tone/間の解析）＋ 終了後に30項目×5段階で採点",
    "新規アーキテクチャ（ストリーミングLLM＋音声＋リアルタイム採点）",
], gap=0.42, dot=ORANGE, size=12.5)
text(s, 0.95, 6.4, 11.6, 0.4, [[("価値：", 12.5, ORANGE, True),("指導者不在でも反復練習、緊張せず壁打ち、早期戦力化", 12.5, GRAY, False)]])

# ============================================================
# Slide 7 — 技術アーキの注意点（実装知見）
# ============================================================
s = slide()
header(s, "LESSONS", "実装知見：避けるべき落とし穴", 6)
items = [
    ("Vercel Serverless(60秒)はロープレに不向き", "バッチ採点ですらタイムアウト多発。リアルタイム音声対話には常駐/エッジ/専用ランタイムが必要。", RED),
    ("1アプリに全部詰めると破綻する", "wheelsUpは機能を足すたびに『分かりにくい』FB。接客版は『1業界・1チェックリスト・1お手本』で別アプリに切る。", ORANGE),
    ("採点はバックグラウンド処理にする", "同期処理は規模が増えると必ずタイムアウト。キュー/Cronで回し、UIは結果表示に徹する。", BLUE),
    ("『印象』でなく『観察可能な行動』で採点", "『良い接客だった』でなく『必須ワードを言った/言わない』。客観性・標準化の肝。", GREEN),
]
y = 1.5
for (h, b, col) in items:
    box(s, 0.6, y, 12.1, 1.2, fill=LIGHT, round_=True); box(s,0.6,y,0.14,1.2,fill=col)
    text(s, 0.95, y+0.16, 11.5, 0.5, [[(h, 15, DARK, True)]])
    text(s, 0.95, y+0.66, 11.6, 0.5, [[(b, 12.5, GRAY, False)]])
    y += 1.32

# ============================================================
# Slide 8 — 推奨アクション
# ============================================================
s = slide()
header(s, "NEXT ACTION", "推奨アクション", 7)
acts = [
    ("1", "フェーズ1（接客採点PoC）を別アプリとして設計", "既存資産で最短に市場の追い風を捉える", BLUE),
    ("2", "接客チェックリストを現場と共同定義", "『正解がある研修』の正解を固める", GREEN),
    ("3", "トップ販売員の接客をお手本データ化", "採点基準の現場実証", ORANGE),
    ("4", "フェーズ2（ロープレ）の技術設計・コスト試算", "投資判断の材料", RED),
]
y = 1.55
for (n, h, b, col) in acts:
    box(s, 0.6, y, 12.1, 1.05, fill=WHITE, line=RGBColor(0xE3,0xE7,0xEC), line_w=1.2, round_=True)
    box(s, 0.78, y+0.2, 0.65, 0.65, fill=col, round_=True)
    text(s, 0.78, y+0.26, 0.65, 0.55, [[(n, 22, WHITE, True)]], align=PP_ALIGN.CENTER)
    text(s, 1.7, y+0.16, 8.0, 0.5, [[(h, 15.5, DARK, True)]])
    text(s, 1.7, y+0.62, 8.0, 0.4, [[(b, 12, GRAY, False)]])
    y += 1.18
box(s, 0.6, y+0.02, 12.1, 0.72, fill=NAVY, round_=True)
text(s, 0.9, y+0.1, 11.5, 0.6,
     [[("結論：今は『作る』より『どこに張るか』を決めるフェーズ。フェーズ1から小さく実証するのが最も勝率が高い。", 14, WHITE, True)]],
     anchor=MSO_ANCHOR.MIDDLE)

out = "/home/user/wheelsUp_Sales/docs/接客AI採点ロープレ事業提案_バイタリティデザイン.pptx"
prs.save(out)
print("SAVED:", out, "slides:", len(prs.slides._sldIdLst))
